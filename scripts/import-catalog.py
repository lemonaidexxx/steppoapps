"""Read-only XLSX import. Usage: python scripts/import-catalog.py path/to/catalog.xlsx"""
import hashlib
import json
import pathlib
import sys
import openpyxl

root = pathlib.Path(__file__).resolve().parents[1]
book = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
sheet = book['2 Unique Offerings']
assert list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))) == ['Modality', 'City', 'Training Center', 'Course Name', 'Hours', 'Slots', 'Course Status']
def stable(prefix, parts):
    return prefix + hashlib.sha256(json.dumps(parts, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()[:20]
# Correct display spelling after hashing original source values so existing IDs remain stable.
name_corrections = {"Janapeses Language A1 Level": "Japanese Language A1 Level"}
rows = []
for index, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
    if not any(v is not None for v in row):
        continue
    mode, city, center, name, hours, slots, status = row
    assert all(isinstance(v, str) and v.strip() for v in [mode, city, center, name])
    assert isinstance(hours, (int, float)) and hours > 0
    rows.append(dict(id=stable('off_', list(row[:6])), courseId=stable('crs_', [name]), courseName=name_corrections.get(name, name), modality='Not specified' if mode=='Not Specified' else mode, hours=hours, institution=center, city=city, selectable=True, sourceRow=index, sourceStatus=status or ''))
assert len(rows) == 139 and len({r['id'] for r in rows}) == 139
assert len({r['courseId'] for r in rows}) == 53
(root/'data').mkdir(exist_ok=True)
(root/'data/catalog.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
print(f'Imported {len(rows)} offerings and 53 course titles.')
